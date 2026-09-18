import io
import os
import docx
import openpyxl
import pypdf
from app.core.logging_config import logger
from app.ingestion.utils import read_file_bytes_safely


def parse_pdf(filepath: str) -> list[dict]:
    """
    Parses PDF file page by page using pypdf.
    Returns a list of structured element dicts preserving page numbers and headings.
    """
    filename = os.path.basename(filepath)
    elements = []
    try:
        logger.info(f"Parsing PDF '{filename}'...")
        pdf_bytes = read_file_bytes_safely(filepath)
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        total_pages = len(reader.pages)
        logger.info(f"Parsed PDF '{filename}' ({total_pages} pages).")

        current_heading = "General"

        for page_idx, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text()
                if not page_text or not page_text.strip():
                    continue

                lines = [line.strip() for line in page_text.splitlines() if line.strip()]
                for line in lines:
                    # Detect potential section headings (short lines, uppercase, or title-like)
                    if len(line) < 80 and (line.isupper() or line.startswith("Section") or line.startswith("Chapter")):
                        current_heading = line

                    elements.append({
                        "text": line,
                        "page": page_idx,
                        "section": current_heading,
                        "title": filename,
                        "type": "paragraph"
                    })
            except Exception as page_err:
                logger.error(f"Error parsing page {page_idx} of '{filename}': {page_err}")
                continue

    except Exception as e:
        logger.error(f"Failed to open PDF document '{filename}': {e}")

    return elements


def parse_xlsx(filepath: str) -> list[dict]:
    """
    Parses Excel worksheets row-by-row with openpyxl.
    Converts rows into key-value structured text strings preserving headers.
    """
    filename = os.path.basename(filepath)
    elements = []
    try:
        logger.info(f"Parsing Excel file '{filename}'...")
        excel_bytes = read_file_bytes_safely(filepath)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # Identify headers from first non-empty row
            header_row_idx = 0
            headers = []
            for idx, r in enumerate(rows):
                non_empty = [str(val).strip() for val in r if val is not None and str(val).strip()]
                if non_empty:
                    header_row_idx = idx
                    headers = [str(val).strip() if val is not None else f"Column_{i+1}" for i, val in enumerate(r)]
                    break

            if not headers:
                continue

            # Process data rows
            for r_idx, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
                row_parts = []
                for h_name, cell_val in zip(headers, row):
                    if cell_val is not None and str(cell_val).strip():
                        row_parts.append(f"{h_name}: {str(cell_val).strip()}")

                if row_parts:
                    formatted_row = f"[Sheet: {sheet_name} | Row {r_idx}]\n" + "\n".join(row_parts)
                    elements.append({
                        "text": formatted_row,
                        "sheet": sheet_name,
                        "row_start": r_idx,
                        "row_end": r_idx,
                        "section": f"Sheet: {sheet_name}",
                        "title": filename,
                        "type": "table_row"
                    })

    except Exception as e:
        logger.error(f"Failed to parse Excel file '{filename}': {e}")

    return elements


def parse_docx(filepath: str) -> list[dict]:
    """
    Parses Word document paragraphs, headings, bullet lists, numbered steps, and tables.
    Formated table rows with header context.
    """
    filename = os.path.basename(filepath)
    elements = []
    try:
        logger.info(f"Parsing Word document '{filename}'...")
        docx_bytes = read_file_bytes_safely(filepath)
        doc = docx.Document(io.BytesIO(docx_bytes))
        current_heading = "General"

        # Paragraphs
        for para in doc.paragraphs:
            val = para.text.strip()
            if not val:
                continue

            style_name = para.style.name.lower() if para.style else ""
            if "heading" in style_name or (para.runs and any(run.bold for run in para.runs if run.text.strip())):
                current_heading = val
                elements.append({
                    "text": val,
                    "section": current_heading,
                    "title": filename,
                    "type": "heading"
                })
            else:
                elements.append({
                    "text": val,
                    "section": current_heading,
                    "title": filename,
                    "type": "paragraph"
                })

        # Tables
        for t_idx, table in enumerate(doc.tables, start=1):
            if not table.rows:
                continue

            # Read table headers from row 0
            header_cells = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells]
            headers = [h if h else f"Col_{i+1}" for i, h in enumerate(header_cells)]

            for r_idx, row in enumerate(table.rows[1:], start=2):
                cell_parts = []
                for h_name, cell in zip(headers, row.cells):
                    cell_text = cell.text.strip()
                    if cell_text:
                        clean_lines = [l.strip() for l in cell_text.splitlines() if l.strip()]
                        cell_parts.append(f"{h_name}: {', '.join(clean_lines)}")

                # Deduplicate consecutive cell strings from merged cells
                unique_parts = []
                for part in cell_parts:
                    if not unique_parts or unique_parts[-1] != part:
                        unique_parts.append(part)

                if unique_parts:
                    formatted_table_row = f"[Table {t_idx} | Row {r_idx}]\n" + "\n".join(unique_parts)
                    elements.append({
                        "text": formatted_table_row,
                        "section": current_heading,
                        "title": filename,
                        "type": "table_row"
                    })

    except Exception as e:
        logger.error(f"Failed to parse Word document '{filename}': {e}")

    return elements

